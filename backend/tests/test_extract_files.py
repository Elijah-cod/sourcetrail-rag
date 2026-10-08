import io
from pathlib import Path

import docx
import pytest

from app.ingest import ExtractionError, extract_file
from app.ingest.models import clean_text, guess_year

FIXTURES = Path(__file__).parent / "fixtures"


def test_pdf_gives_one_section_per_page():
    doc = extract_file("annual-report-2024.pdf", (FIXTURES / "sample.pdf").read_bytes())
    assert doc.source_type == "pdf"
    assert doc.title == "Sample Annual Report"          # from PDF metadata
    assert [s.page for s in doc.sections] == [1, 2]
    assert "declined in 2024" in doc.sections[1].text
    assert doc.published_year == 2024                    # guessed from the filename


def test_markdown_splits_on_headings_and_ignores_code_blocks():
    md = b"# Guide\nIntro text.\n\n## Install\nRun it.\n```\n# not a heading\n```\n## Usage\nUse it."
    doc = extract_file("guide.md", md)
    assert doc.title == "Guide"
    assert [s.heading for s in doc.sections] == ["Guide", "Install", "Usage"]
    assert "# not a heading" in doc.sections[1].text


def test_docx_groups_paragraphs_under_headings():
    document = docx.Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.add_heading("Details", level=1)
    document.add_paragraph("More detail.")
    buffer = io.BytesIO()
    document.save(buffer)

    doc = extract_file("notes.docx", buffer.getvalue())
    assert [s.heading for s in doc.sections] == ["Overview", "Details"]
    assert doc.sections[0].text == "First paragraph.\n\nSecond paragraph."


def test_csv_makes_one_section_per_row():
    data = b"product,revenue,year\nWidget,100,2023\nGadget,,2024\n"
    doc = extract_file("sales.csv", data)
    assert [s.row for s in doc.sections] == [1, 2]
    assert doc.sections[0].text == "product: Widget; revenue: 100; year: 2023"
    assert doc.sections[1].text == "product: Gadget; year: 2024"   # empty cells skipped


def test_txt_falls_back_to_latin1():
    doc = extract_file("old_notes.txt", "Caf\xe9 menu".encode("latin-1"))
    assert doc.title == "old notes"
    assert doc.sections[0].text == "Café menu"


def test_unsupported_type_is_rejected():
    with pytest.raises(ExtractionError, match="Unsupported"):
        extract_file("picture.png", b"\x89PNG")


def test_empty_file_is_rejected():
    with pytest.raises(ExtractionError, match="No text"):
        extract_file("empty.txt", b"   \n\n ")


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ExtractionError, match="Could not read PDF"):
        extract_file("broken.pdf", b"this is not a pdf")


def test_helpers():
    assert guess_year("report_2019_final.pdf") == 2019
    assert guess_year("v12345.pdf") is None
    assert clean_text("a  b \r\n\n\n\nc\x00") == "a b\n\nc"


# ---------- PDF structure (tests/fixtures/structured.pdf, see make_structured_pdf.py) ----------

def structured():
    return extract_file("handbook (2).pdf", (FIXTURES / "structured.pdf").read_bytes())


def test_pdf_placeholder_metadata_title_is_ignored():
    # The metadata says "untitled"; the big heading on page 1 is the real title.
    assert structured().title == "Aurora Operations Handbook"


def test_pdf_headings_lists_and_paragraphs_are_kept():
    page1 = structured().sections[0].text
    assert page1.startswith("# Aurora Operations Handbook\n\n## 1. Service Availability\n\n")
    assert "\n\n### Maintenance windows\n\n" in page1
    assert ("- Customers receive notice at least 72 hours ahead.\n"
            "- Emergency fixes may happen at any time.\n") in page1
    # Wrapped lines are joined back into one paragraph.
    assert "Planned maintenance is excluded from the calculation when customers are notified in advance." in page1


def test_pdf_running_header_and_page_numbers_are_dropped():
    for section in structured().sections:
        assert "Aurora Operations Handbook" not in section.text.replace("# Aurora Operations Handbook", "")
        assert not any(line.strip().isdigit() for line in section.text.splitlines())


def test_pdf_heading_carries_over_page_breaks():
    pages = structured().sections
    assert [p.heading for p in pages] == [None, "Maintenance windows", "2. Incident Response"]


def test_file_name_fallback_title_is_readable():
    from app.ingest.files import _stem
    assert _stem("BSR+110+-+Theological+Writing+Skills (8).pdf") == "BSR 110 - Theological Writing Skills"
    assert _stem("annual_report-2024.pdf") == "annual report 2024"


def test_hyphenated_line_breaks_are_rejoined():
    from app.ingest.pdf_layout import _join
    assert _join("planned main-", "tenance") == "planned maintenance"
    assert _join("state-of-the-", "Art") == "state-of-the- Art"      # capital: a real hyphen, keep it


# ---------- CSV header validation (SourceTrail test report, 8 Oct 2026) ----------
# A headerless CSV was accepted: its first data row became the field names, so the index held
# "RG-X1: RG-X2; North pier: East Lock; 12: 7; Ready: Overdue" and answers cited it confidently.

@pytest.mark.parametrize("data, message", [
    (b"RG-X1,North pier,12,Ready\nRG-X2,East Lock,7,Overdue\n", "'12' as a column name"),   # the report's file
    (b"RG-X1,North pier,Ready\nRG-X2,East Lock,Overdue\n", "'RG-X1' as a column name"),     # no numeric column
    (b"North pier,Ready\nEast Lock,Ready\nSouth quay,Overdue\n", "doesn't seem to have a header"),
    (b"asset,location\n", "at least one data row"),
    (b"asset,,status\nRG-X1,North pier,Ready\n", "empty column name"),
    (b"asset,Asset\nRG-X1,RG-X2\n", "more than once"),
    (b"2024-01-05,location\nRG-X1,North pier\n", "'2024-01-05' as a column name"),
    ((b"x" * 150 + b",location\nRG-X1,North pier\n"), "longer than 100"),
    (b"", "no header row"),
])
def test_csv_without_a_real_header_is_rejected(data, message):
    with pytest.raises(ExtractionError, match=message):
        extract_file("assets.csv", data)


def test_csv_with_header_still_accepted():
    data = b"asset,location,days,status,note\nRG-X1,North pier,12,Ready,\"oiled, tested\"\nRG-X2,East Lock,7,Overdue,\n"
    doc = extract_file("assets.csv", data)
    assert doc.sections[0].text == "asset: RG-X1; location: North pier; days: 12; status: Ready; note: oiled, tested"
    assert doc.sections[1].text == "asset: RG-X2; location: East Lock; days: 7; status: Overdue"


def test_csv_trailing_comma_column_is_ignored_not_rejected():
    doc = extract_file("assets.csv", b"name,status,\nRamp,Ready,\nGate,Ready,\n")
    assert [s.text for s in doc.sections] == ["name: Ramp; status: Ready", "name: Gate; status: Ready"]


# ---------- TXT sections (SourceTrail test report, 8 Oct 2026) ----------
# The whole TXT file was one section, so neighbouring conditional rules shared a chunk and the
# answer to a warning question dropped "extinguish the main lantern" and added a fuel reorder.

STANDING_ORDERS = """HARBOUR LIGHT STANDING ORDERS
=============================

Keeper on duty: Anaïs Vézin. Relief keeper: Noël Étoile.

Lamp Room

The main lantern is lit at sunset and extinguished at sunrise.
The lens is polished every morning.

If the fog bell sounds for more than ten minutes, extinguish the main lantern and switch to the auxiliary beam.

When the fuel tank falls below one quarter, reorder fuel from Côte Rêverie depot.

General warnings:

- Never leave the lamp room unattended.
- Report broken glass immediately.

2. Succession

If the keeper is absent, the relief keeper takes command.
""".encode("latin-1")


def test_txt_splits_at_headings_and_isolates_conditional_rules():
    sections = extract_file("standing_orders.txt", STANDING_ORDERS).sections
    by_text = {s.text: s.heading for s in sections}
    fog = ("If the fog bell sounds for more than ten minutes, extinguish the main lantern "
           "and switch to the auxiliary beam.")
    fuel = "When the fuel tank falls below one quarter, reorder fuel from Côte Rêverie depot."
    assert by_text[fog] == "Lamp Room"                  # the rule and its actions, alone together
    assert by_text[fuel] == "Lamp Room"
    assert by_text["Keeper on duty: Anaïs Vézin. Relief keeper: Noël Étoile."] == "HARBOUR LIGHT STANDING ORDERS"
    # A "...:" intro stays with the list it introduces.
    assert by_text["General warnings:\n- Never leave the lamp room unattended.\n- Report broken glass immediately."] == "Lamp Room"
    assert by_text["If the keeper is absent, the relief keeper takes command."] == "2. Succession"


def test_txt_rules_never_share_a_chunk():
    from app.chunking import chunk_document
    chunks = chunk_document(extract_file("standing_orders.txt", STANDING_ORDERS), "orders")
    assert not any("fog bell" in c.text and "fuel tank" in c.text for c in chunks)


def test_txt_plain_paragraphs_are_grouped_not_one_chunk_each():
    text = "\n\n".join(f"Note {i}: the boat shed holds oars and spare rope." for i in range(4))
    assert len(extract_file("notes.txt", text.encode()).sections) == 1


# ---------- DOCX tables, headers and footers (SourceTrail test report, 8 Oct 2026) ----------

def make_docx_with_table_header_footer() -> bytes:
    document = docx.Document()
    document.add_heading("Stores", level=1)
    document.add_paragraph("Inventory is checked weekly.")
    table = document.add_table(rows=3, cols=3)
    for r, values in enumerate([["Item", "Qty", "Station"], ["Lamp oil", "4", "North pier"],
                                ["Wick", "12", "East Lock"]]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    document.add_paragraph("Shortfalls go to the harbour master.")
    document.sections[0].header.paragraphs[0].text = "Harbour Board internal: code word HERON"
    document.sections[0].footer.paragraphs[0].text = "Issued by the Lighthouse Directorate"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_tables_are_extracted_row_by_row_in_place():
    sections = extract_file("stores.docx", make_docx_with_table_header_footer()).sections
    assert [s.heading for s in sections[:3]] == ["Stores", "Stores", "Stores"]
    assert sections[0].text == "Inventory is checked weekly."
    assert sections[1].text == ("Table 1 columns: Item | Qty | Station\n"
                                "Table 1, row 1: Item: Lamp oil; Qty: 4; Station: North pier\n"
                                "Table 1, row 2: Item: Wick; Qty: 12; Station: East Lock")
    assert sections[2].text == "Shortfalls go to the harbour master."


def test_docx_headers_and_footers_are_labelled_sections():
    sections = extract_file("stores.docx", make_docx_with_table_header_footer()).sections
    labelled = {s.heading: s.text for s in sections if s.heading in ("Page header", "Page footer")}
    assert labelled == {"Page header": "Harbour Board internal: code word HERON",
                        "Page footer": "Issued by the Lighthouse Directorate"}
